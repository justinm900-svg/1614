"""Builds '1614 Quote Calculator.xlsx' (CURRENT 2026-09-24). Formulas only; mirrors quote_calculator.py."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

wb = Workbook()
F = "Arial"
blue = Font(name=F, color="0000FF"); black = Font(name=F); bold = Font(name=F, bold=True)
hdr = Font(name=F, bold=True, size=13); yellow = PatternFill("solid", fgColor="FFFF00")
grey = PatternFill("solid", fgColor="EDEDED"); thin = Side(style="thin", color="999999")
money = '$#,##0.00;($#,##0.00);-'

# ---------------- Pricing sheet ----------------
p = wb.active; p.title = "Pricing"
p["A1"] = "1614 Home Co. — Pricing (source of truth, launch brief 2026-09-24)"; p["A1"].font = hdr
rows = [
 ("Standard — up to 3 baths", 799), ("Standard — 4 baths", 899), ("Standard — 5 baths", 999),
 ("Founding — up to 3 baths", 699), ("Founding — 4 baths", 799), ("Founding — 5 baths", 899),
 ("Custom review at bathrooms >=", 6), ("Additional HVAC system (per year)", 150),
 ("Additional water heater (per year)", 40), ("Member Setup base", 199),
 ("Additional standard leak sensor (beyond the 4 included in setup)", 35), ("Specialty / probe sensor", 49),
 ("Water Defense standalone setup", 249), ("Membership conversion credit", 50),
 ("Monthly factor", 1.08), ("Months in term", 12), ("Founding cap (homes, TOTAL)", 25),
 ("Custom review at sq ft >=", 5000), ("Water Defense conversion window (days)", 30),
]
for i, (k, v) in enumerate(rows, start=3):
    p.cell(i, 1, k).font = black; c = p.cell(i, 2, v); c.font = blue
    if isinstance(v, (int, float)) and v not in (6, 12, 25, 5000, 30, 1.08): c.number_format = money
p["B21"].comment = Comment("Rule (Justin 2026-09-24): homeowner who completed the $249 Water Defense Setup and joins membership within 30 days pays no $199 Member Setup, is charged only newly approved sensors/equipment, and gets a separate $50 credit on the first-year membership price.", "1614")
p["A25"] = "Sensor rule (Justin 2026-09-24): Member Setup $199 includes the hub + first 4 standard leak sensors. Enter only ADDITIONAL standard sensors beyond 4 (= max(total at home - 4, 0)). Specialty sensors are $49 each; a specialty sensor in one of the first 4 spots is $49 only."; p["A25"].font = Font(name=F, italic=True, color="666666")
p["A24"] = "Tier rule (Justin 2026-09-24): each full AND each half bathroom counts as one bathroom toward the tier. 2 full + 2 half = 4-bath tier."; p["A24"].font = Font(name=F, italic=True, color="666666")
p["A23"] = "Blue = input from launch brief (Justin, 2026-09-24). Do not create new pricing here."; p["A23"].font = Font(name=F, italic=True, color="666666")
p.column_dimensions["A"].width = 44; p.column_dimensions["B"].width = 14
N = {  # named refs into Pricing
 "s3":"Pricing!$B$3","s4":"Pricing!$B$4","s5":"Pricing!$B$5","f3":"Pricing!$B$6","f4":"Pricing!$B$7","f5":"Pricing!$B$8",
 "cb":"Pricing!$B$9","hv":"Pricing!$B$10","wh":"Pricing!$B$11","su":"Pricing!$B$12","std":"Pricing!$B$13","spc":"Pricing!$B$14",
 "wd":"Pricing!$B$15","cr":"Pricing!$B$16","fac":"Pricing!$B$17","mo":"Pricing!$B$18","cap":"Pricing!$B$19","sq":"Pricing!$B$20","win":"Pricing!$B$21"}

# ---------------- Quote sheet ----------------
q = wb.create_sheet("Quote")
q["A1"] = "1614 Home Co. — Quote Calculator (CURRENT 2026-09-24)"; q["A1"].font = hdr
q["A2"] = "Keep water where it belongs.  Fill YELLOW cells only. Everything else is a formula."; q["A2"].font = Font(name=F, italic=True)
inputs = [
 ("Customer / address label", "TEST — Fake Customer (do not invoice)", None),
 ("Full bathrooms", 2, None), ("Half bathrooms", 1, None), ("HVAC systems", 2, None), ("Water heaters", 1, None),
 ("Approx. square feet (blank if unknown)", 2400, None), ("Sump pump? (yes/no/blank)", "yes", None),
 ("Quote a membership? (1=yes, 0=Water Defense only)", 1, None), ("Founding? (1/0)", 1, None),
 ("Additional Standard Sensors Beyond Included 4 (count; = max(total - 4, 0))", 3, None), ("Specialty / Probe Sensors (count)", 1, None),
 ("Water Defense setup? (1/0)", 0, None), ("Water Defense conversion? (1/0) — setup completed within 30 days", 0, None),
 ("Boiler? (1/0)", 0, None), ("Well / pressure tank? (1/0)", 0, None), ("Unusual complexity? (1/0)", 0, None),
 ("Sales tax rate (fraction, e.g. 0.08; blank = not applied)", None, None),
 ("Tax applies to (none / setup / all)", "none", None),
]
q["A4"] = "INPUTS"; q["A4"].font = bold
for i, (k, v, _) in enumerate(inputs, start=5):
    q.cell(i, 1, k).font = black; c = q.cell(i, 2, v); c.font = blue; c.fill = yellow
# input cell map
I = {k: f"$B${i}" for i, (k, _, _) in enumerate(inputs, start=5)}
full, half, hvac, wh, sqft, sump = I["Full bathrooms"], I["Half bathrooms"], I["HVAC systems"], I["Water heaters"], I["Approx. square feet (blank if unknown)"], I["Sump pump? (yes/no/blank)"]
mem, fnd, nstd, nspc, wdf, crd = I["Quote a membership? (1=yes, 0=Water Defense only)"], I["Founding? (1/0)"], I["Additional Standard Sensors Beyond Included 4 (count; = max(total - 4, 0))"], I["Specialty / Probe Sensors (count)"], I["Water Defense setup? (1/0)"], I["Water Defense conversion? (1/0) — setup completed within 30 days"]
boil, well, cplx, trate, tbase = I["Boiler? (1/0)"], I["Well / pressure tank? (1/0)"], I["Unusual complexity? (1/0)"], I["Sales tax rate (fraction, e.g. 0.08; blank = not applied)"], I["Tax applies to (none / setup / all)"]

r = 5 + len(inputs) + 1
q.cell(r, 1, "CALCULATION").font = bold; r += 1
calc = []
def add(label, formula, fmt=money, b=False):
    global r
    q.cell(r, 1, label).font = bold if b else black
    c = q.cell(r, 2, formula); c.font = bold if b else black
    if fmt: c.number_format = fmt
    calc.append((label, f"$B${r}")); r += 1
    return f"$B${r-1}"
wcount = add("Bathroom count (full + half, each counts as one)", f"={full}+{half}", "0")
tier = add("Bathroom tier (3/4/5, 0 = custom review)", f"=IF({wcount}>={N['cb']},0,IF({wcount}<=3,3,{wcount}))", "0")
custom = add("Custom review required? (1/0)", f"=IF(OR({tier}=0,AND({sqft}<>\"\",{sqft}>={N['sq']})),1,0)", "0")
active = f"AND({mem}=1,{custom}=0)"
conv = f"AND({crd}=1,{mem}=1,{custom}=0)"
base = add("Membership base", f"=IF({active},IF({fnd}=1,CHOOSE({tier}-2,{N['f3']},{N['f4']},{N['f5']}),CHOOSE({tier}-2,{N['s3']},{N['s4']},{N['s5']})),0)")
hvt = add("Additional HVAC adder", f"=IF({active},MAX(0,{hvac}-1)*{N['hv']},0)")
wht = add("Additional water heater adder", f"=IF({active},MAX(0,{wh}-1)*{N['wh']},0)")
crt = add("Water Defense conversion credit (first-year membership)", f"=IF({conv},-{N['cr']},0)")
annual = add("ANNUAL TOTAL, first year (membership + adders + conversion credit)", f"=ROUND({base}+{hvt}+{wht}+{crt},2)", b=True)
sub = add("Member Setup base — includes hub + first 4 standard sensors (waived on Water Defense conversion)", f"=IF(AND({active},NOT({conv})),{N['su']},0)")
stdt = add("Additional standard sensors beyond included 4 (× $35)", f"={nstd}*{N['std']}")
spct = add("Specialty / probe sensors (× $49)", f"={nspc}*{N['spc']}")
wdt = add("Water Defense setup (not charged again on conversion)", f"=IF(AND({wdf}=1,NOT({conv})),{N['wd']},0)")
setup = add("SETUP TOTAL (one time, upfront; conversion = new sensors/equipment only)", f"=ROUND({sub}+{stdt}+{spct}+{wdt},2)", b=True)
monthly = add("MONTHLY PAYMENT (annual × 1.08 ÷ 12, rounded to cents)", f"=IF({annual}>0,ROUND({annual}*{N['fac']}/{N['mo']},2),0)", b=True)
add("Monthly total over 12-month term", f"=ROUND({monthly}*{N['mo']},2)")
tax = add("Sales tax PLACEHOLDER (not a taxability determination)", f"=IF(OR({trate}=\"\",{tbase}=\"none\"),0,ROUND(IF({tbase}=\"setup\",{setup},{annual}+{setup})*{trate},2))")
add("PAY ANNUALLY: due at signup (annual + setup + tax placeholder)", f"=ROUND({annual}+{setup}+{tax},2)", b=True)
r += 1
q.cell(r, 1, "FLAGS").font = bold; r += 1
flags = [
 f'=IF({tier}=0,"CUSTOM REVIEW: 6+ bathrooms — do not quote from table.","")',
 f'=IF(AND({sqft}<>"",{sqft}>={N["sq"]}),"CUSTOM REVIEW: >= 5,000 sq ft.","")',
 f'=IF({boil}=1,"FLAG: Boiler present — confirm scope before quoting.","")',
 f'=IF({well}=1,"FLAG: Well / pressure tank — confirm scope before quoting.","")',
 f'=IF({cplx}=1,"FLAG: Unusual complexity — review before quoting.","")',
 f'=IF({sump}="","INFO: Sump pump status not confirmed — ask on call.","")',
 f'=IF({fnd}=1,"FOUNDING: verify seat available (first "&{N["cap"]}&" homes TOTAL across launch area).","")',
 f'=IF(AND({crd}=1,NOT({active})),"NOTE: Water Defense conversion applies only when joining a membership; not applied.","")',
 f'=IF({conv},"CONVERSION: confirm Water Defense Setup completed within "&{N["win"]}&" days. No Member Setup fee; only newly approved sensors/equipment charged; $50 credit on first-year membership.","")',
 f'=IF(AND({conv},{wdf}=1),"NOTE: Water Defense Setup already purchased; $249 not charged again.","")',
 f'=IF(OR({trate}="",{tbase}="none"),"Sales tax: TBD — taxability pending CPA confirmation; rate is address-based (Ohio The Finder).","Sales tax PLACEHOLDER applied — verify before invoicing.")',
]
for fml in flags:
    q.cell(r, 1, fml).font = Font(name=F, color="8A1F1F"); r += 1
r += 1
q.cell(r, 1, "Legend: yellow/blue = fill in; black = formula; Pricing sheet holds every rate. Monthly rounding is half-up via ROUND(). Setup is paid upfront; 12-month term.").font = Font(name=F, italic=True, color="666666")
q.column_dimensions["A"].width = 62; q.column_dimensions["B"].width = 40

# ---------------- Test cases sheet (expected values from quote_calculator.py) ----------------
t = wb.create_sheet("Test Cases")
t["A1"] = "Expected results from quote_calculator.py — paste inputs into Quote sheet to spot-check"; t["A1"].font = hdr
hdrs = ["Case", "Full", "Half", "HVAC", "WH", "Founding", "Addl std sensors (beyond 4)", "Spec sensors", "Water Def", "Conversion", "Expected annual", "Expected setup", "Expected monthly"]
for j, h in enumerate(hdrs, 1): c = t.cell(3, j, h); c.font = bold; c.fill = grey
cases = [
 ("Std 2/1 baths", 2,1,1,1,0,0,0,0,0, 799, 199, 71.91),
 ("Std 4 baths", 4,0,1,1,0,0,0,0,0, 899, 199, 80.91),
 ("Std 5 baths", 5,0,1,1,0,0,0,0,0, 999, 199, 89.91),
 ("Founding 2/1 + 2 HVAC + 3 ADDITIONAL std + 1 spec (Scenario A)", 2,1,2,1,1,3,1,0,0, 849, 353, 76.41),
 ("Founding 2/1 + 2 HVAC + 3 TOTAL std (0 additional) + 1 spec", 2,1,2,1,1,0,1,0,0, 849, 248, 76.41),
 ("Founding 3/1 (tier 4) + 2 WH", 3,1,1,2,1,0,0,0,0, 839, 199, 75.51),
 ("Std 2 baths + Water Defense", 2,0,1,1,0,0,0,1,0, 799, 448, 71.91),
 ("Std 2/2 baths (4 bathrooms -> 4-bath tier)", 2,2,1,1,0,0,0,0,0, 899, 199, 80.91),
 ("Std 2/1 + WD conversion + 1 new std sensor", 2,1,1,1,0,1,0,0,1, 749, 35, 67.41),
 ("Founding 3/1 + 2 HVAC + WD conversion", 3,1,2,1,1,0,0,0,1, 899, 0, 80.91),
 ("WD conversion with WD box ticked (no double charge)", 2,0,1,1,0,0,0,1,1, 749, 0, 67.41),
 ("6 baths -> custom review", 6,0,1,1,0,0,0,0,0, 0, 0, 0),
]
for i, row in enumerate(cases, start=4):
    for j, v in enumerate(row, 1):
        c = t.cell(i, j, v); c.font = blue if j > 1 else black
        if j >= 11: c.number_format = money
for col in "ABCDEFGHIJKLM": t.column_dimensions[col].width = 14
t.column_dimensions["A"].width = 40
wb.save("1614 Quote Calculator.xlsx"); print("saved")
